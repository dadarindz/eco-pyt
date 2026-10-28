"""Génère un compte rendu .docx (et .pdf) simple à partir d'un fichier texte
au format Markdown réduit, puis supprime toutes les métadonnées.

Usage : python3 generer_cr.py chemin/vers/cr.md [--pdf]

Syntaxe reconnue dans cr.md :
    # Titre du TP
    > Nom Prénom | Classe | Date          (ligne d'en-tête, optionnelle)
    ## Partie
    ### Sous-partie
    - élément de liste
    ```  ... ```                          (bloc de commandes)
    ![Légende](captures/fichier.png)      (capture d'écran)
    texte normal
"""
import io
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from pypdf import PdfReader, PdfWriter

POLICE = "Calibri"
TAILLE = 11

CORE_VIDE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
    'xmlns:dcmitype="http://purl.org/dc/dcmitype/" '
    'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"/>'
)
APP_VIDE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"/>'
)


def appliquer_police(run, taille=TAILLE, gras=False):
    run.font.name = POLICE
    run.font.size = Pt(taille)
    run.font.bold = gras
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), POLICE)


def ajouter_texte(par, texte, taille=TAILLE, gras=False):
    # **gras** et `code` en ligne : même police, seul le gras change
    for morceau in re.split(r"(\*\*[^*]+\*\*)", texte):
        if not morceau:
            continue
        est_gras = morceau.startswith("**") and morceau.endswith("**")
        contenu = morceau[2:-2] if est_gras else morceau
        appliquer_police(par.add_run(contenu.replace("`", "")), taille, gras or est_gras)


def image_sans_exif(chemin):
    img = Image.open(chemin)
    propre = Image.frombytes(img.mode, img.size, img.tobytes())
    tampon = io.BytesIO()
    propre.save(tampon, format="PNG")
    tampon.seek(0)
    return tampon


def construire(source: Path, sortie: Path):
    doc = Document()
    for nom in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet", "Caption"):
        style = doc.styles[nom]
        style.font.name = POLICE
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), POLICE)
    doc.styles["Normal"].font.size = Pt(TAILLE)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.5)

    lignes = source.read_text(encoding="utf-8").splitlines()
    num_figure = 0
    i = 0
    while i < len(lignes):
        ligne = lignes[i].rstrip()
        if ligne.startswith("```"):
            i += 1
            while i < len(lignes) and not lignes[i].startswith("```"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                p.paragraph_format.space_after = Pt(0)
                appliquer_police(p.add_run(lignes[i]))
                i += 1
            doc.add_paragraph()
        elif ligne.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ajouter_texte(p, ligne[2:], 18, True)
        elif ligne.startswith("> "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ajouter_texte(p, ligne[2:])
        elif ligne.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            ajouter_texte(p, ligne[3:], 14, True)
        elif ligne.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            ajouter_texte(p, ligne[4:], 12, True)
        elif ligne.startswith("- "):
            ajouter_texte(doc.add_paragraph(style="List Bullet"), ligne[2:])
        elif m := re.match(r"!\[(.*)\]\((.+)\)", ligne):
            num_figure += 1
            doc.add_picture(image_sans_exif(source.parent / m.group(2)), width=Cm(15))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ajouter_texte(p, f"Figure {num_figure} : {m.group(1)}", 10)
        elif ligne:
            ajouter_texte(doc.add_paragraph(), ligne)
        i += 1

    doc.save(sortie)
    nettoyer_docx(sortie)


def nettoyer_docx(chemin: Path):
    """Réécrit le .docx sans auteur, dates, application, ni propriétés perso."""
    tmp = chemin.with_suffix(".tmp")
    with zipfile.ZipFile(chemin) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename == "docProps/custom.xml":
                continue
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = CORE_VIDE.encode()
            elif item.filename == "docProps/app.xml":
                data = APP_VIDE.encode()
            elif item.filename in ("[Content_Types].xml", "_rels/.rels"):
                data = re.sub(rb"<[^>]*custom[-.]?(?:properties|xml)[^>]*/>", b"", data)
            info = zipfile.ZipInfo(item.filename, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(info, data)
    shutil.move(tmp, chemin)


def exporter_pdf(docx: Path):
    with tempfile.TemporaryDirectory() as d:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", d, str(docx)],
                       check=True, capture_output=True)
        lecteur = PdfReader(Path(d) / (docx.stem + ".pdf"))
        ecrivain = PdfWriter()
        for page in lecteur.pages:
            ecrivain.add_page(page)
        ecrivain.add_metadata({"/Producer": "", "/Creator": "", "/Author": "", "/Title": ""})
        with open(docx.with_suffix(".pdf"), "wb") as f:
            ecrivain.write(f)


if __name__ == "__main__":
    source = Path(sys.argv[1]).resolve()
    sortie = source.with_suffix(".docx")
    construire(source, sortie)
    print("Compte rendu :", sortie)
    if "--pdf" in sys.argv:
        exporter_pdf(sortie)
        print("PDF :", sortie.with_suffix(".pdf"))
